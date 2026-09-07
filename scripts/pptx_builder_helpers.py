import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# ─── COLOR PALETTE ───────────────────────────────────────────────────────────
C_DARK_BG       = RGBColor(3, 18, 36)      # Deep Sapphire Dark
C_CARD_BG       = RGBColor(8, 30, 58)      # Card dark background
C_CARD_BORDER   = RGBColor(20, 60, 105)    # Card subtle border
C_ACCENT_CYAN   = RGBColor(56, 189, 248)   # Bright Cyan
C_ACCENT_BLUE   = RGBColor(2, 132, 199)    # Royal Blue
C_ACCENT_AMBER  = RGBColor(245, 158, 11)   # Warning / Aging Amber
C_ACCENT_GREEN  = RGBColor(16, 185, 129)   # Success Green
C_ACCENT_RED    = RGBColor(239, 68, 68)    # Critical Red
C_TEXT_WHITE    = RGBColor(255, 255, 255)  # Headings / primary
C_TEXT_MUTED    = RGBColor(148, 163, 184)  # Subtitles / secondary
C_TEXT_BODY     = RGBColor(226, 232, 240)  # Body text
C_PILL_BG       = RGBColor(14, 45, 82)     # Badge pill bg

FONT_NAME = "Segoe UI"
FONT_BOLD = "Segoe UI"

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs

def add_solid_background(slide, color=C_DARK_BG):
    bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = color
    bg_shape.line.fill.background()
    return bg_shape

def add_header(slide, kicker, title, category="SAHA YÖNETİM PANELİ (SYP)"):
    # Header container
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.733), Inches(1.1))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    # Category & Kicker
    p0 = tf.paragraphs[0]
    p0.text = f"{category.upper()}  •  {kicker.upper()}"
    p0.font.name = FONT_BOLD
    p0.font.size = Pt(9.5)
    p0.font.bold = True
    p0.font.color.rgb = C_ACCENT_CYAN
    p0.space_after = Pt(3)

    # Main Title
    p1 = tf.add_paragraph()
    p1.text = title
    p1.font.name = FONT_BOLD
    p1.font.size = Pt(22)
    p1.font.bold = True
    p1.font.color.rgb = C_TEXT_WHITE

def add_takeaway(slide, text, top=1.65):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(top), Inches(11.733), Inches(0.52))
    shape.fill.solid()
    shape.fill.fore_color.rgb = C_PILL_BG
    shape.line.color.rgb = C_ACCENT_BLUE
    shape.line.width = Pt(1)

    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_right = Inches(0.2)
    tf.margin_top = Inches(0.08)
    p = tf.paragraphs[0]
    p.text = f"💡 YÖNETİCİ VURGUSU:  {text}"
    p.font.name = FONT_BOLD
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = C_TEXT_WHITE

def add_card(slide, left, top, width, height, title=None, icon=None, bg=C_CARD_BG, border=C_CARD_BORDER):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    card.fill.solid()
    card.fill.fore_color.rgb = bg
    card.line.color.rgb = border
    card.line.width = Pt(1.2)

    if title:
        tb = slide.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.18), Inches(width - 0.5), Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        full_title = f"{icon}  {title}" if icon else title
        p.text = full_title
        p.font.name = FONT_BOLD
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = C_ACCENT_CYAN

    return card

def add_card_content(slide, left, top, width, height, items):
    tb = slide.shapes.add_textbox(Inches(left + 0.25), Inches(top), Inches(width - 0.5), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.text = item
        p.font.name = FONT_NAME
        p.font.size = Pt(10)
        p.font.color.rgb = C_TEXT_BODY
        p.space_after = Pt(6)

def add_kpi_card(slide, left, top, width, height, icon, value, label, subtext="", accent=C_ACCENT_CYAN):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    card.fill.solid()
    card.fill.fore_color.rgb = C_CARD_BG
    card.line.color.rgb = accent
    card.line.width = Pt(1.5)

    tb = slide.shapes.add_textbox(Inches(left + 0.15), Inches(top + 0.15), Inches(width - 0.3), Inches(height - 0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    p0 = tf.paragraphs[0]
    p0.text = f"{icon}  {value}"
    p0.font.name = FONT_BOLD
    p0.font.size = Pt(20)
    p0.font.bold = True
    p0.font.color.rgb = accent
    p0.space_after = Pt(2)

    p1 = tf.add_paragraph()
    p1.text = label
    p1.font.name = FONT_BOLD
    p1.font.size = Pt(11)
    p1.font.bold = True
    p1.font.color.rgb = C_TEXT_WHITE
    p1.space_after = Pt(3)

    if subtext:
        p2 = tf.add_paragraph()
        p2.text = subtext
        p2.font.name = FONT_NAME
        p2.font.size = Pt(9)
        p2.font.color.rgb = C_TEXT_MUTED

def add_footer(slide, current, total):
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.733), Inches(0.35))
    tf = tb.text_frame
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = f"Saha Yönetim Paneli (SYP)  •  Meydan Yönetimi Birimi  •  Sayfa {current} / {total}"
    p.font.name = FONT_NAME
    p.font.size = Pt(8.5)
    p.font.color.rgb = C_TEXT_MUTED

print("PPTX Helper library initialized successfully.")
